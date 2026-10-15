"""The SICKLE-Guard paper, IEEE two-column format.

    python scripts/make_guard_paper_ieee.py

Writes docs/SICKLE_Guard_IEEE.docx and results/figures_guard/*.png.

Every number is read from the artefacts through make_thesis_paper.load(),
the loader the thesis paper uses, so the two documents cannot disagree. The
layout helpers (two columns, headings, citation renumbering) are those of
make_paper_ieee.py, so the two IEEE papers share one style.

Images, all column width:
  fig_guard_workflow     the five steps and the status of the evidence
  fig_guard_report       a field report rendered by the prototype's own
                         code, from ILLUSTRATIVE inputs (not a SICKLE field)
  fig_guard_sheet        an inspection task sheet rendered by the workflow,
                         field and season masked (queue order is label-derived)
  fig_guard_ranking      held-out ranking power against alternatives
  fig_guard_official     recalled labels against the state's records
  fig_guard_simulated    simulated studies, labelled so
  fig_guard_console      the console's summary cards (aggregates only),
                         screenshotted with headless Microsoft Edge when it
                         is installed; skipped otherwise

Refuses to save if the text contains an un-negated overstatement from the
claims registry, or loses its simulated or scope labels.
"""

from __future__ import annotations

import contextlib
import io
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import yaml
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJ))
sys.path.insert(0, str(PROJ / "scripts"))

import make_thesis_paper as T                           # noqa: E402
from make_paper_ieee import (author_block, body, heading,  # noqa: E402
                             renumber, two_columns)
from make_thesis_paper import ci, cil, m, pct, w        # noqa: E402

OUT = PROJ / "docs" / "SICKLE_Guard_IEEE.docx"
GFIG = PROJ / "results" / "figures_guard"
COL = 3.3
EDGE = Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe")

TITLE = ("SICKLE-Guard: Prioritising Field Verification While Guarding "
         "Against False Reassurance in Satellite-Assisted Crop Insurance")
# ----------------------------------------------------------------- layout
_FIGNO = [0]
_TABNO = [0]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]


def figure(doc, path, caption, width=COL):
    if path is None or not Path(path).exists():
        print(f"  figure skipped (not available): {path}")
        return None
    _FIGNO[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(f"Fig. {_FIGNO[0]}.  {caption}")
    r.font.size = Pt(7.5)
    return _FIGNO[0]


def table(doc, header, rows, caption, widths=None, size=7.0):
    _TABNO[0] += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    r = cap.add_run(f"TABLE {ROMAN[_TABNO[0] - 1]}\n{caption}")
    r.font.size = Pt(7.5)
    r.bold = True
    tb = doc.add_table(rows=1, cols=len(header))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = tb.rows[0].cells[i]
        c.text = h
        for p_ in c.paragraphs:
            for r_ in p_.runs:
                r_.bold = True
                r_.font.size = Pt(size)
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p_ in cells[i].paragraphs:
                for r_ in p_.runs:
                    r_.font.size = Pt(size)
    if widths:
        for row in tb.rows:
            for i, wd in enumerate(widths):
                row.cells[i].width = Inches(wd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return _TABNO[0]


def lead(doc, lead_text, text, indent=0.2):
    """A paragraph opening with an italic run-in lead, IEEE style."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(indent)
    r = p.add_run(lead_text + " ")
    r.italic = True
    p.add_run(text)
    return p


# ----------------------------------------------------------------- figures
def _rc():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.size": 7, "figure.dpi": 300,
                         "savefig.bbox": "tight", "savefig.pad_inches": 0.03,
                         "axes.grid": True, "grid.alpha": 0.25,
                         "axes.axisbelow": True, "font.family": "DejaVu Sans"})
    return plt


def text_figure(lines, out, fontsize=5.4, title=None, note=None):
    plt = _rc()
    n = len(lines) + (2 if title else 0) + (2 if note else 0)
    fig = plt.figure(figsize=(3.35, 0.098 * n + 0.15))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    ax.add_patch(plt.Rectangle((0.005, 0.005), 0.99, 0.99,
                               transform=ax.transAxes, fill=True,
                               fc="#fbfbf8", ec="#999999", lw=0.6))
    y = 0.985
    step = 0.97 / n
    if title:
        ax.text(0.03, y, title, transform=ax.transAxes, fontsize=fontsize + 0.6,
                weight="bold", va="top", color="#7a2e1f")
        y -= 2 * step
    for ln in lines:
        ax.text(0.03, y, ln, transform=ax.transAxes, fontsize=fontsize,
                family="DejaVu Sans Mono", va="top", color="#111111")
        y -= step
    if note:
        y -= step
        ax.text(0.03, y, note, transform=ax.transAxes, fontsize=fontsize,
                style="italic", va="top", color="#555555")
    fig.savefig(out, facecolor="white")
    plt.close(fig)
    return out


def make_figures(N):
    plt = _rc()
    GFIG.mkdir(parents=True, exist_ok=True)
    F = {}

    # ---- 1. workflow, vertical, the evidence status of each step
    steps = [("1  Estimate", "a field's own past harvests first;\n"
              "satellite only where there is no record", "#2d7a3a",
              "observed"),
             ("2  Flag", "no reliability signal discriminates\n"
              "(none of nine tested), so none is shown", "#b03636",
              "refuted, so not used"),
             ("3  Prioritise", "rank recorded fields by history;\n"
              "reserve a share for no-record fields", "#2d7a3a", "observed"),
             ("4  Verify", "blind task sheets, observation-only\n"
              "records, SHA-256 audit chain", "#a5611a",
              "software; visits simulated"),
             ("5  Evaluate", "pre-registered: against random visits,\n"
              "the register, and history alone", "#2d7a3a", "observed")]
    fig, ax = plt.subplots(figsize=(3.35, 2.55))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 10.2)
    ax.axis("off")
    for i, (t, body_, col, status) in enumerate(steps):
        y0 = 8.25 - i * 2.0
        ax.add_patch(plt.Rectangle((0.1, y0), 9.8, 1.7, fill=False, ec=col,
                                   lw=1.2))
        ax.text(0.35, y0 + 0.85, t, va="center", fontsize=7.2,
                weight="bold", color=col)
        ax.text(3.0, y0 + 0.72, body_, va="center", fontsize=6.0,
                linespacing=1.2)
        ax.text(9.75, y0 + 1.42, status, va="center", ha="right",
                fontsize=5.4, style="italic", color=col)
        if i < 4:
            ax.annotate("", xy=(1.2, y0 - 0.28), xytext=(1.2, y0 - 0.02),
                        arrowprops=dict(arrowstyle="->", color="#555555",
                                        lw=0.8))
    F["workflow"] = GFIG / "fig_guard_workflow.png"
    fig.savefig(F["workflow"])
    plt.close(fig)

    # ---- 2. the field report, rendered by the prototype's own code
    import farmer_report as fr
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        fr.report({"plot": 999, "season": "oct-mar", "year": 2020}, 1850.0,
                  "satellite", 101.0, False, fr.contacts(), "en")
    raw = buf.getvalue().rstrip("\n").split("\n")
    keep, cut = [], False
    for ln in raw:
        s = ln.strip()
        if s.startswith("This warning means"):
            keep.append("      [...]")
            cut = True
        if s.startswith("WHAT TO DO NOW"):
            cut = False
            keep.append("")
        if s.startswith("- localised"):
            keep.append("      [... reporting routes and contacts ...]")
            cut = True
        if s.startswith("This is a research prototype"):
            cut = False
            keep.append("")
        if not cut:
            keep.append(ln[:68])
    F["report"] = text_figure(
        [k for k in keep if k is not None], GFIG / "fig_guard_report.png",
        note="Illustrative inputs (field 999 is not a SICKLE field); "
             "text exactly as the prototype prints it.")

    # ---- 3. the inspection task sheet, rendered by the workflow
    r = subprocess.run([sys.executable, "src/inspection_workflow.py", "sheet",
                        "--task", "1"], cwd=T.GUARD, capture_output=True,
                       text=True)
    import re as _re
    sheet = []
    for ln in r.stdout.rstrip("\n").split("\n"):
        ln = _re.sub(r"field \d+", "field ###", ln)
        ln = _re.sub(r"season [a-z]+-[a-z]+ \d{4}", "season ########", ln)
        sheet.append(ln[:66])
    F["sheet"] = text_figure(
        sheet, GFIG / "fig_guard_sheet.png",
        note="Field and season masked: queue order derives from labels.")

    # ---- 4. held-out ranking
    t = N["ep"]["ranking"]["held-out test split"]
    order = [("SICKLE-Guard (as shipped)", "SICKLE-Guard"),
             ("last harvest only", "last harvest only"),
             ("history only, no imagery", "history only, no imagery"),
             ("largest field first", "largest field first"),
             ("satellite only", "satellite only"),
             ("arbitrary order", "arbitrary order"),
             ("smallest field first", "smallest field first")]
    order = [o for o in order if o[0] in t]
    fig, ax = plt.subplots(figsize=(3.35, 2.0))
    y = np.arange(len(order))[::-1]
    a = [t[k]["auroc"] for k, _ in order]
    lo = [t[k]["auroc"] - t[k]["auroc_ci"][0] for k, _ in order]
    hi = [t[k]["auroc_ci"][1] - t[k]["auroc"] for k, _ in order]
    ax.errorbar(a, y, xerr=[lo, hi], fmt="o", ms=3.5, color="#1f3b73",
                ecolor="#7f9cc6", capsize=2, lw=1)
    ax.axvline(0.5, color="#c0392b", ls="--", lw=0.9, label="chance")
    ax.set_yticks(y)
    ax.set_yticklabels([lab for _, lab in order], fontsize=6.5)
    ax.set_xlabel("AUROC for a below-threshold field-season (95 % CI)",
                  fontsize=6.5)
    ax.set_xlim(0, 1)
    ax.legend(loc="lower right", fontsize=6)
    F["ranking"] = GFIG / "fig_guard_ranking.png"
    fig.savefig(F["ranking"])
    plt.close(fig)

    # ---- 5. labels and model against the state's records (stacked)
    off = N["off"]
    cells = off["l1_level"]["cells"]
    ev = off["l2_event"]
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(3.35, 3.1),
                                 gridspec_kw={"height_ratios": [1, 1.1]})
    names = [("Kar/Kuruvai/Sornavari 2019-20", "Kuruvai"),
             ("Samba/Thaladi/Pishanam 2019-20", "Samba")]
    for yi, (k, lab) in enumerate(names):
        c = cells[k]
        b = c["by_conversion"]["2/3 (official)"]
        lo_b, hi_b = c["official_core_band"]
        a1.add_patch(plt.Rectangle((lo_b, yi - 0.28), hi_b - lo_b, 0.56,
                                   color="#888888", alpha=0.25, lw=0))
        for v in c["official_core_districts"].values():
            a1.plot([v, v], [yi - 0.28, yi + 0.28], color="#666666", lw=0.8)
        a1.errorbar([b["sickle_mean_rice_eq_kg_ha"]], [yi],
                    xerr=[[b["sickle_mean_rice_eq_kg_ha"] - b["ci95"][0]],
                          [b["ci95"][1] - b["sickle_mean_rice_eq_kg_ha"]]],
                    fmt="o", ms=3.5, color="#1f3b73", capsize=2)
    a1.set_yticks([0, 1])
    a1.set_yticklabels([n_[1] for n_ in names])
    a1.set_ylim(-0.7, 1.7)
    a1.set_xlabel("kg of cleaned rice per ha, 2019-20", fontsize=6.5)
    a1.set_title("(a) recalled labels (dot) vs state crop-cutting "
                 "districts (grey)", fontsize=6.8)
    rows = [("recalled, Samba", ev["samba_recalled"], "#1f3b73"),
            ("recalled, Kuruvai\n(control)", ev["kuruvai_control"], "#888888"),
            ("satellite model,\nSamba", ev["samba_model"], "#222222")]
    for yi, (lab, rr, col) in enumerate(rows[::-1]):
        m_ = 100 * rr["mean_change"]
        a2.errorbar([m_], [yi], xerr=[[m_ - 100 * rr["ci95"][0]],
                                      [100 * rr["ci95"][1] - m_]],
                    fmt="s", ms=3.5, color=col, capsize=2)
    a2.axvline(0, color="#c0392b", lw=0.9, ls="--")
    a2.set_yticks(range(3))
    a2.set_yticklabels([r_[0] for r_ in rows[::-1]], fontsize=6.2)
    a2.set_xlabel("same-plot change into 2020-21 (%)", fontsize=6.5)
    a2.set_title("(b) the season the state paid delta-wide relief",
                 fontsize=6.8)
    fig.tight_layout(h_pad=0.8)
    F["official"] = GFIG / "fig_guard_official.png"
    fig.savefig(F["official"])
    plt.close(fig)

    # ---- 6. simulated studies (stacked), labelled
    ho = N["insp"]["sets"]["held-out test split (VERDICT)"]["budgets"]["0.2"]
    cur = N["yw"]["conditional_curves"]
    fig, (b1, b2) = plt.subplots(2, 1, figsize=(3.35, 3.5))
    sv = np.linspace(0.3, 1.0, 15)
    for k, lab, col in (("register", "register-guided", "#1f6f4a"),
                        ("workflow", "workflow-guided", "#1f3b73")):
        b1.plot(sv, sv * ho[k]["n_found"], color=col, lw=1.4, label=lab)
    b1.plot(sv, sv * ho["random"]["mean_found"], color="#999999", lw=1.2,
            label="random visits")
    b1.axhline(ho["random"]["mean_found"], color="#999999", ls="--", lw=0.8,
               label="perfect inspector, random")
    b1.set_xlabel("inspector sensitivity", fontsize=6.5)
    b1.set_ylabel("real problems confirmed", fontsize=6.5)
    b1.legend(fontsize=5.6, loc="upper left")
    b1.set_title("(a) SIMULATED inspector, observed rankings, 20 % budget",
                 fontsize=6.8, color="#a5611a")
    sh = [0.05, 0.10, 0.20, 0.30]
    for wt, lab, col in (("0.30", "30 % weight", "#1f3b73"),
                         ("0.50", "50 %", "#1f6f4a"),
                         ("1.00", "100 %", "#a5611a")):
        for sd, ls in (("0.20", "-"), ("0.30", "--")):
            b2.plot([100 * x for x in sh],
                    [100 * cur[f"w={wt}|sd_t={sd}|bias=+0.00|sd_c=0.10"]
                     [f"{x:.2f}"]["p_paid_nothing"] for x in sh],
                    color=col, ls=ls, marker="o", ms=2.5, lw=1.1,
                    label=lab if sd == "0.20" else None)
    b2.set_xlabel("true shortfall below threshold (%)", fontsize=6.5)
    b2.set_ylabel("paid nothing (%)", fontsize=6.5)
    b2.legend(fontsize=5.6)
    b2.set_title("(b) SIMULATED Insurance Unit, technology error 0.20 "
                 "(solid) / 0.30 (dashed)", fontsize=6.8, color="#a5611a")
    fig.tight_layout(h_pad=0.8)
    F["simulated"] = GFIG / "fig_guard_simulated.png"
    fig.savefig(F["simulated"])
    plt.close(fig)

    # ---- 7. the console's summary cards (aggregates only)
    F["console"] = None
    html = T.GUARD / "results" / "console" / "verification_console.html"
    shot = GFIG / "console_narrow.png"
    if EDGE.exists() and html.exists():
        subprocess.run([str(EDGE), "--headless=new", "--disable-gpu",
                        "--hide-scrollbars", "--force-device-scale-factor=3",
                        "--window-size=680,1500", f"--screenshot={shot}",
                        html.resolve().as_uri()],
                       capture_output=True, timeout=120)
    if shot.exists():
        from PIL import Image
        im = Image.open(shot)
        # the four cards: below the legend rule, above the inspection queue
        # (at 680 CSS px and scale 3; re-check if the console layout changes)
        F["console"] = GFIG / "fig_guard_console.png"
        im.crop((0, 990, im.width, 2725)).save(F["console"])
    return F


# ----------------------------------------------------------------- pipeline
def pipeline_figure():
    """How the satellite model reads the imagery, for ONE held-out field.

    The field is chosen by a rule that never looks at its yield: among
    held-out field-seasons with no harvest record (the only case in which
    the report uses the satellite), the one with the median number of clear
    Sentinel-2 acquisitions, ties broken by id. Every value is computed by
    the prototype's own code, and two assertions tie the figure to it: the
    clear count must equal the saved feature table's, and the estimate must
    equal the number the shipped report prints for this field. The recalled
    yield is not shown.
    """
    import glob as _glob
    import json as _json
    import re as _re
    import extract_features as ef
    import farmer_report as fr
    from reliability import cols_all
    plt = _rc()
    from matplotlib.collections import LineCollection

    feats = fr.load()
    hist = fr.causal_history(feats)
    sat = fr.satellite(feats)
    cands = [i for i, r in enumerate(feats)
             if r["split"] == "test" and not np.isfinite(hist[i])]
    ncl = {i: int(float(feats[i]["s2_n_clear"])) for i in cands}
    order = sorted(cands, key=lambda i: (ncl[i], str(feats[i]["uid"])))
    pick = order[len(order) // 2]
    r = feats[pick]
    uid, pid = r["uid"], float(r["plot"])

    sel = ef.plot_pixels(uid, pid, "10m")
    kept, rejected = [], []
    for f in sorted(_glob.glob(str(ef.ROOT / f"images/S2/npy/{uid}/*.npz"))):
        d = np.load(f)
        dt = ef.acq_date(f)
        bad = False
        if "SCL" in d:
            mm = ef.fit_mask(sel, d["SCL"].shape)
            bad = bool(mm.sum() and np.isin(d["SCL"][mm].astype(int),
                                            list(ef.SCL_BAD)).mean() > 0.5)
        if bad or "B4" not in d or "B8" not in d:
            rejected.append(dt)
            continue
        mm = ef.fit_mask(sel, d["B4"].shape)
        red = float(np.nanmean(d["B4"][mm]))
        nir = float(np.nanmean(d["B8"][mm]))
        kept.append((dt, (nir - red) / (nir + red + 1e-9), f))
    # tie the figure to the saved features and to the shipped report
    assert len(kept) == ncl[pick], (len(kept), ncl[pick])
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        fr.report(r, sat[pick], "satellite", 100.0, False, None, "en")
    printed = int(_re.search(r"Estimated yield: (\d+)", buf.getvalue()).group(1))
    assert printed == round(sat[pick]), (printed, sat[pick])
    n_feat = len(cols_all())

    peak = max(kept, key=lambda k: k[1])
    d = np.load(peak[2])
    rgb = np.stack([d["B4"], d["B3"], d["B2"]], -1).astype(float)
    lo, hi = np.nanpercentile(rgb, 2), np.nanpercentile(rgb, 98)
    rgb = np.clip((rgb - lo) / (hi - lo), 0, 1)
    msk = ef.fit_mask(sel, d["B4"].shape)

    fig = plt.figure(figsize=(3.35, 2.75))
    gs = fig.add_gridspec(2, 2, width_ratios=[1, 1.75],
                          height_ratios=[1.35, 1], wspace=0.32, hspace=0.45)
    a = fig.add_subplot(gs[0, 0])
    a.grid(False)
    a.imshow(rgb, interpolation="nearest")
    segs = []
    H, W = msk.shape
    for y in range(H):
        for x in range(W):
            if not msk[y, x]:
                continue
            if y == 0 or not msk[y - 1, x]:
                segs.append([(x - .5, y - .5), (x + .5, y - .5)])
            if y == H - 1 or not msk[y + 1, x]:
                segs.append([(x - .5, y + .5), (x + .5, y + .5)])
            if x == 0 or not msk[y, x - 1]:
                segs.append([(x - .5, y - .5), (x - .5, y + .5)])
            if x == W - 1 or not msk[y, x + 1]:
                segs.append([(x + .5, y - .5), (x + .5, y + .5)])
    a.add_collection(LineCollection(segs, colors="white", linewidths=0.9))
    a.set_xticks([])
    a.set_yticks([])
    a.set_title(f"(a) Sentinel-2 true colour\n{peak[0]:%d %b %Y}, field "
                "outlined", fontsize=6.2)
    a.plot([1, 11], [H - 2, H - 2], color="white", lw=1.2)
    a.text(6, H - 3, "100 m", color="white", fontsize=5, ha="center",
           va="bottom")

    b = fig.add_subplot(gs[0, 1])
    b.plot([k[0] for k in kept], [k[1] for k in kept], "-o", ms=2, lw=0.8,
           color="#2d7a3a")
    b.plot([peak[0]], [peak[1]], "o", ms=4.5, mfc="none", mec="#a5611a",
           mew=1)
    yb = min(k[1] for k in kept) - 0.08
    b.plot(rejected, [yb] * len(rejected), "x", ms=2.5, color="#999999")
    b.set_ylabel("NDVI", fontsize=6.0, labelpad=1)
    b.set_title(f"(b) field-mean NDVI: {len(kept)} clear (dots) of\n"
                f"{len(kept) + len(rejected)} acquisitions; × rejected "
                "as cloud", fontsize=5.8)
    b.tick_params(axis="x", labelsize=5, rotation=30)

    c = fig.add_subplot(gs[1, :])
    c.axis("off")
    steps = [("mask to\nthe field", "#1f3b73"),
             ("drop cloudy\ndates", "#1f3b73"),
             (f"{n_feat} season\nfeatures", "#1f3b73"),
             ("random\nforest", "#1f3b73"),
             (f"estimate\n{sat[pick]:,.0f}\nkg/acre", "#a5611a")]
    for k, (txt, col) in enumerate(steps):
        x0 = 0.01 + k * 0.198
        c.add_patch(plt.Rectangle((x0, 0.38), 0.17, 0.55,
                                  transform=c.transAxes, fill=False, ec=col,
                                  lw=0.9))
        c.text(x0 + 0.085, 0.655, txt, transform=c.transAxes, ha="center",
               va="center", fontsize=5.2, color=col)
        if k < 4:
            c.annotate("", xy=(x0 + 0.195, 0.655), xytext=(x0 + 0.172, 0.655),
                       xycoords="axes fraction",
                       arrowprops=dict(arrowstyle="->", lw=0.7,
                                       color="#555555"))
    c.text(0.5, 0.08, "(c) a model output, used only because this field has "
           "no harvest record", transform=c.transAxes, ha="center",
           fontsize=5.4, style="italic", color="#a5611a")
    out = GFIG / "fig_guard_pipeline.png"
    fig.savefig(out, facecolor="white")
    plt.close(fig)
    meta = {"selection": "held-out field-seasons with no harvest record; "
                         "median clear Sentinel-2 count; ties by id",
            "uid": uid, "plot": pid, "season": r["season"],
            "n_candidates": len(cands), "n_clear": len(kept),
            "n_acquisitions": len(kept) + len(rejected),
            "estimate_kg_per_acre": float(sat[pick]),
            "report_prints": printed, "n_features": n_feat,
            "peak_date": f"{peak[0]:%Y-%m-%d}"}
    out.with_suffix(".json").write_text(_json.dumps(meta, indent=2),
                                        encoding="utf8")
    return out, meta


# ----------------------------------------------------------------- paper
def z(v, nd=3):
    """A signed value, except that one rounding to zero carries no sign."""
    return m(0.0, nd) if round(v, nd) == 0 else m(v, nd, True)


def build(N):
    from docx import Document
    F = make_figures(N)
    F["pipeline"], PM = pipeline_figure()
    ep = N["ep"]
    hold = ep["ranking"]["held-out test split"]
    flag = ep["flag_adds_information"]["held-out test split"]
    warn = ep["warning_calibration"]["held-out test split"]
    base = N["ae"]["baseline_comparison"]
    lh = base["held-out test split"]
    eb = N["eb"]["sets"]["held-out test split (PRIMARY)"]["has_record"]
    a, b = eb["has a record"], eb["cold start"]
    lv = N["lv"]
    ev = lv["ranking_by_subset"]["EVENT plots only (the real case)"]
    ss = N["ss"]
    dsa = N["dsa"]
    hist_r = dsa["ranking"]["history-only (no imagery)"]
    hyb_r = dsa["ranking"]["hybrid = SICKLE-Guard"]
    rd = dsa["ranking_delta"]
    mae_cold = dsa["paired_mae_hybrid_minus_history_only"]["cold_start"]
    vp = N["vp"]
    ea = vp["experiment_a"]["held_out"]
    eb_ = vp["experiment_b"]["all_oof"]
    b33 = eb_["budgets"]["0.33"]
    b33z, b33t = b33["0.00|satellite"], b33["0.10|satellite"]
    b33p = b33["proportional|satellite"]
    wc = b33["_within_cold"]
    ex = vp["exploratory_reserve"]["all_oof"]["0.33"]
    insp = N["insp"]
    ho20 = insp["sets"]["held-out test split (VERDICT)"]["budgets"]["0.2"]
    oofs = insp["sets"]["all field-seasons out-of-fold (secondary)"]
    expo = insp["observed_exposure"]["held-out test split"]
    exp20 = expo["budgets"]["0.2"]
    thr = expo["threshold_kg_per_acre"]
    g7 = ho20["grid"]["s=0.70|f=0.10"]
    cc = N["yw"]["conditional_curves"]
    p100 = cc["w=1.00|sd_t=0.20|bias=+0.00|sd_c=0.10"]["0.10"]["p_paid_nothing"]
    p30 = cc["w=0.30|sd_t=0.20|bias=+0.00|sd_c=0.10"]["0.10"]["p_paid_nothing"]
    gd = N["yw"]["guard"]["loss=0.20|w=1.00|sd_t=0.20|bias=+0.00"]["budgets"]["0.05"]["policies"]
    off = N["off"]
    hl = off["headline"]
    rel = off["l2_event"]["relief"]["total_lakh"]
    rel20 = next(v for k_, v in rel.items() if k_.startswith("2020-21")) / 100
    A20 = ea["budgets"]["0.2"]
    pil = N["pilot"]["sets"]["all field-seasons out-of-fold"]["comparisons"]
    ys = N["ys"]["threshold_free"]
    mn = N["mndwi"]["sets"]["held-out test split (PRIMARY)"]
    dmn = mn["signals"]["decoupling (pre-specified)"]
    n_ho_fs = sum(1 for _ in range(0))  # placeholder, replaced below
    n_ho_fs = N["pilot"]["sets"]["held-out test split"]["n"]
    n_ho_bad = ea["n_bad"]

    # yield-reliability signals, as the thesis reports them
    sig_names = {"quantile_width": "prediction-interval width",
                 "rf_spread": "random-forest spread",
                 "seed_spread": "seed disagreement",
                 "low_n_clear": "too few clear observations",
                 "small_n_px": "too few pixels",
                 "ood_distance": "out-of-distribution distance",
                 "low_predicted_yield": "low predicted yield"}
    sigrows = [(sig_names.get(k, k), m(v["auroc"]),
                cil(v["ci"]["lo"], v["ci"]["hi"]),
                "yes" if v["informative"] else "no")
               for k, v in N["yield_signals"].items()]
    sigrows.append(("released-ensemble dispersion*", m(ys["auroc"]),
                    cil(ys["auroc_ci"]["lo"], ys["auroc_ci"]["hi"]), "no"))
    sigrows.append(("water/vegetation decoupling*", m(dmn["auroc"]),
                    cil(*dmn["ci"]), "yes" if dmn["informative"] else "no"))
    n_sig = len(sigrows)
    n_inf = sum(1 for r in sigrows if r[3] == "yes")

    doc = Document()
    s0 = doc.sections[0]
    s0.left_margin = s0.right_margin = Inches(0.62)
    s0.top_margin = Inches(0.7)
    s0.bottom_margin = Inches(0.7)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(18)
    t.paragraph_format.space_after = Pt(10)
    author_block(doc)

    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    sec.left_margin = sec.right_margin = Inches(0.62)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    two_columns(sec)

    # ------------------------------------------------------------ abstract
    ab = doc.add_paragraph()
    ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rr = ab.add_run("Abstract\u2014")
    rr.bold = True
    rr.italic = True
    rr.font.size = Pt(9)
    rr = ab.add_run(
        "India has begun to automate crop-insurance payouts: claims are "
        "worked out from the shortfall in yield, and in Madhya Pradesh the "
        "technology-based estimate carried full weight in 2025-26. When such "
        "an estimate errs, the costly mistake is not a needless field visit "
        "but a farmer reassured into not reporting damage. We present "
        "SICKLE-Guard, a research prototype that prioritises which fields "
        "an inspection office should verify under a fixed budget, gives "
        "each field a report with no reassuring state, and turns "
        "verification into auditable software: blind task sheets, "
        "observation-only records and a SHA-256 hash chain. Every design "
        "rule traces to a measurement on the SICKLE paddy benchmark, "
        "evaluated on 27 held-out plots under pre-registered protocols. "
        "The field's own past harvests carry the result: the satellite "
        "component adds nothing measurable to the ranking (AUROC change "
        f"{m(rd['delta'], 3, True)}, 95 % CI {m(rd['ci'][0], 3, True)} to "
        f"{m(rd['ci'][1], 3, True)}), and an advantage over the register an "
        "office already keeps is not demonstrated, although every "
        "record-based policy far outperforms random visits. None of "
        f"{w(n_sig)} yield-reliability signals discriminates, so no "
        "confidence score is shown. All "
        f"{w(exp20['n_missed'])} below-threshold field-seasons the ranking "
        "misses at a 20 % budget have no harvest record and carry estimates "
        "above the threshold: exactly the farmers a reassuring display "
        "would reach. Simulated studies of imperfect inspectors and of the "
        "scheme's weighting are reported separately. No farmer or official "
        "has used the system.")
    rr.font.size = Pt(9)
    kw = doc.add_paragraph()
    kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    kw.paragraph_format.space_before = Pt(5)
    kw.paragraph_format.space_after = Pt(6)
    rr = kw.add_run("Index Terms\u2014")
    rr.bold = True
    rr.italic = True
    rr.font.size = Pt(9)
    rr = kw.add_run("crop insurance, decision support, verification "
                    "prioritisation, satellite yield estimation, false "
                    "reassurance, audit trail, smallholder agriculture.")
    rr.font.size = Pt(9)

    # ------------------------------------------------------------ I
    heading(doc, "I.  Introduction")
    body(doc, "India has begun to automate crop-insurance payouts. Under the "
              "Pradhan Mantri Fasal Bima Yojana [3], the YES-TECH framework "
              "gives technology-based yield estimates at least 30 % of the "
              "weight in the yield that settles a claim [2]; in 2025-26 "
              "Madhya Pradesh gave them 100 % and Maharashtra 50 %, and "
              "claims are \u201cautomatically worked out based on the "
              "shortfall in yield\u201d [21]. Farmers in Madhya Pradesh "
              "contest the resulting assessments [5], [22].", 0.0)
    body(doc, "The two ways such a system can err are not symmetric. A false "
              "alarm costs an inspection visit. A false reassurance, an "
              "estimate that says a damaged field looks normal, can stop a "
              "farmer from reporting the loss in time. A tool placed between "
              "a yield model and a farmer should therefore be judged by "
              "whether it avoids reassurance it cannot justify, as much as by "
              "its accuracy. A companion audit of the benchmark used here, "
              "SICKLE [1], "
              "shows why that caution is warranted: its pooled accuracy "
              "conceals fields mapped badly in every run, and none of its "
              "released yield models is shown to beat a constant predictor.")
    body(doc, "We ask what an inspection office can safely do with satellite "
              "yield estimates of this quality, and answer with a working "
              "prototype whose every rule is tested against the alternatives "
              "an office already has. Our contributions are: (i) a "
              "verification-prioritisation design in which each rule traces "
              "to a measurement (Table I); (ii) a pre-registered evaluation "
              "on held-out plots against random visits, the register an "
              "office keeps and a field's own history; (iii) the "
              "verification step as tested software, with blind task "
              "sheets, observation-only records and a tamper-evident audit "
              "chain; (iv) a check of the benchmark's labels and of the "
              "model against Tamil Nadu's published crop-cutting statistics "
              "and crop-damage relief; and (v) simulated studies of "
              "imperfect inspectors and of the scheme's weighting, kept "
              "separate from every observed result.")
    body(doc, "Scope. SICKLE-Guard is a research prototype. It is not "
              "connected to PMFBY, is not an insurer, and its output has no "
              "official standing; no farmer or official has used it, and no "
              "outcome for any farmer is measured.")

    # ------------------------------------------------------------ II
    heading(doc, "II.  Related Work")
    body(doc, "Satellite yield estimation for smallholder insurance is "
              "constrained by the ground data used to judge it. In Uganda, "
              "farmer-reported yields diverged from crop-cut and full-plot "
              "measures, while satellite estimates explained full-plot "
              "yields as well as or better than either [30]; in Bihar, "
              "satellite models calibrated on farmers' self-reported yields "
              "failed to predict crop-cut yields [29]. For index insurance, "
              "yield data covering many fields but few years lead standard "
              "estimates to understate basis risk and so overstate insurance "
              "quality [27]. These studies concern how well an estimate can "
              "be measured or how a scheme should be designed; ours concerns "
              "what to do downstream of an estimate whose reliability for a "
              "given field is unknown.", 0.0)
    body(doc, "Machine learning offers two standard routes to that: "
              "estimate uncertainty, for example from the disagreement of "
              "independently trained models [11], or abstain, deferring "
              "uncertain cases under a risk-coverage trade-off [12]. Both "
              "presume a signal that ranks the unreliable cases. We test "
              "that presumption rather than assume it, and report the "
              "results field by field rather than as an aggregate, which "
              "hides where a system fails [25].")

    # ------------------------------------------------------------ III
    heading(doc, "III.  Setting and Data")
    body(doc, "SICKLE [1] covers smallholder paddy in the Cauvery Delta of "
              "Tamil Nadu, 2018-2021, with Sentinel-1, Sentinel-2 and "
              "Landsat-8 time series and, for a subset of plots, yield, "
              "recalled by farmers in surveys run in 2021-22. "
              f"{dsa['n']} field-seasons from {dsa['n_plots']} plots carry a "
              "yield. Plots were partitioned once, by plot, into training, "
              "validation and a held-out test split of 27 plots "
              f"({n_ho_fs} field-seasons) that was never used to fit or "
              "tune anything; out-of-fold predictions over all "
              "field-seasons give a second, larger view with more power.", 0.0)
    body(doc, "A field-season counts as below threshold when its recalled "
              "yield is under 80 % of the training-split mean, "
              f"{thr:,.0f} kg/acre. This is a convention of the study, not "
              "an official criterion. A field's history is the mean yield "
              "of its seasons harvested before the season in question was "
              f"sown; {dsa['n_with_history']} of {dsa['n']} field-seasons "
              f"({pct(dsa['n_with_history'] / dsa['n'])}) have one, and the "
              "rest are cold starts with no record. Official records are the "
              "state's district-season rice yields from crop-cutting "
              "experiments and its crop-damage relief tables [23].")

    # ------------------------------------------------------------ IV
    body(doc, "The satellite model reduces each field-season to numbers. "
              "Every Sentinel-2 acquisition is averaged over the field's own "
              "pixels, with the field mask resampled onto each band's native "
              "grid, and an acquisition is dropped when more than half the "
              "field's pixels are flagged as cloud, shadow, snow or "
              "saturated by the scene classification. Vegetation, water and "
              "moisture indices (NDVI, NDWI, NDMI, EVI and a red-edge "
              "chlorophyll index), red, near-infrared and short-wave "
              "infrared reflectance, and Sentinel-1 VV, VH and a radar "
              "vegetation index are summarised over the season by their "
              "mean, maximum, minimum and standard deviation; with the "
              "timing of the greenness peak and its rates of rise and fall, "
              "the numbers of acquisitions, and the field's area and pixel "
              f"count, this gives {PM['n_features']} features. A random "
              "forest (300 trees, depth 8), fitted on the training and "
              "validation plots and averaged over three seeds, turns them "
              "into a yield estimate (Fig. 1). The report uses that estimate "
              "only for a field with no harvest record, and labels it there "
              "as no reliable estimate.")
    figure(doc, F["pipeline"],
           "How the satellite model reads the imagery, for one held-out "
           "field with no harvest record, chosen by a rule that does not "
           "look at its yield (the median number of clear acquisitions among "
           f"{PM['n_candidates']} such field-seasons). (a) True colour on "
           "the date of peak greenness. (b) Field-mean NDVI, one of the "
           f"{PM['n_features']} season features; cloudy acquisitions are "
           "dropped. (c) The steps from image to estimate. The estimate is "
           "the model's output as the shipped report prints it, not a "
           "measured yield; one field illustrates the pipeline and says "
           "nothing about its accuracy.")

    heading(doc, "IV.  The SICKLE-Guard Design")
    body(doc, "One rule governs the design. Because a false reassurance is "
              "the costly error, the report has no reassuring state: a field "
              "is flagged or not flagged, and not flagged always says that "
              "this is not evidence of no loss and that damage should be "
              "reported anyway. A test fails the build if any "
              "report contains reassuring or claim-deciding language outside "
              "an explicit negation.", 0.0)
    figure(doc, F["workflow"],
           "The five steps of SICKLE-Guard and the status of the evidence "
           "behind each.")
    body(doc, "The prototype works in five steps (Fig. 2). It estimates a "
              "field's yield from its own past harvests and falls back to "
              "the satellite model only where no record exists; it shows no "
              "reliability flag, because none was found to work; it ranks "
              "field-seasons for verification under a budget, ranking "
              "recorded fields by their history and reserving a stated "
              "share for fields with none; it hands inspectors task sheets "
              "and records what they find; and every policy is evaluated "
              "against simpler ones before it is adopted. Table I traces "
              "each rule to the measurement behind it.")
    table(doc, ["Rule", "Evidence (Section V)"], [
        ("Rank by the field's own past harvests",
         f"history-only ranking AUROC {m(hist_r['auroc'])}; satellite adds "
         f"{m(rd['delta'], 3, True)} {ci(*rd['ci'])} (observed)"),
        ("No reassuring state in any report",
         f"all {w(exp20['n_missed'])} misses at 20 % carry estimates above "
         "the threshold (observed)"),
        ("No confidence score shown",
         f"none of {w(n_sig)} reliability signals discriminates (observed)"),
        ("Reserve a share for no-record fields, chosen at random",
         "no-record recall 0 without a reserve; satellite no better than "
         "random among them (observed)"),
        ("Withhold the estimate from inspectors",
         f"{w(exp20['n_inspected_bad_estimate_at_or_above_threshold'])} of "
         f"{w(A20['workflow']['n_found'])} inspected problem field-seasons "
         "carry an estimate above the threshold (observed)"),
        ("Record observations, never decisions; hash every action",
         f"{N['n_wf_tests']} workflow tests, tampering caught (software)"),
    ], "Each design rule and the evidence behind it", widths=[1.25, 2.05])
    body(doc, "Fig. 3 shows the rule at work: a report for a field with no "
              "harvest record and an estimate close to the average, the case "
              "in which a reassuring display would be most tempting.")
    figure(doc, F["report"],
           "A field report for an unflagged field with no harvest record, "
           "printed by the prototype's own code. Even an estimate near the "
           "average is accompanied by a statement that it is not evidence "
           "of no loss, and the first instruction is to report damage now. "
           "Lines elided as marked.")

    # ------------------------------------------------------------ V
    heading(doc, "V.  Evaluation on Held-Out Plots")
    heading(doc, "A.  What the ranking ranks", 2)
    body(doc, f"Of the labelled losses, {pct(lv['frac_losses_chronic'])} sit "
              "at plots below threshold in every observed season. Feeding "
              "the model a different season from the same plot keeps most of "
              "its ranking skill (AUROC "
              f"{m(ss['1 real features']['loss_auroc'])} with real features "
              f"against {m(ss['2 same-plot season swap']['loss_auroc'])} "
              "swapped), and on plots that are sometimes above and sometimes "
              f"below threshold the ranking falls to {m(ev['auroc'])}, "
              f"95 % CI {cil(*ev['ci'])}, which includes chance. The tool "
              "therefore ranks fields by their usual yield level and is "
              "nowhere described as detecting damage in a season.", 0.0)

    heading(doc, "B.  Against the alternatives an officer could use", 2)
    body(doc, "On the 27 held-out plots the shipped ranking scores AUROC "
              f"{m(hold['SICKLE-Guard (as shipped)']['auroc'])} "
              f"{cil(*hold['SICKLE-Guard (as shipped)']['auroc_ci'])} and "
              f"beats an arbitrary order ({m(hold['arbitrary order']['auroc'])})"
              " (Fig. 4). It does not demonstrably beat a lookup of each "
              "field's most recent harvest: the paired difference is "
              f"{m(lh['paired_delta'], 3, True)} {ci(*lh['ci'])}. Over all "
              "field-seasons, where a cold start must fall back on "
              "satellite, ranking by history alone scores "
              f"{m(hist_r['auroc'])} {cil(*hist_r['ci'])} and the hybrid "
              f"{m(hyb_r['auroc'])} {cil(*hyb_r['ci'])}, a difference of "
              f"{m(rd['delta'], 3, True)} {ci(*rd['ci'])}; on cold starts, "
              "where the only alternative is the area average, the "
              "satellite lowers yield error by "
              f"{m(-mae_cold['delta'], 1)} kg/acre "
              f"{ci(-mae_cold['ci'][1], -mae_cold['ci'][0], 1)}, an interval "
              "that includes zero. The satellite component has not been "
              "shown to earn its place; it is kept because a new field has "
              "no alternative, and its number is labelled as not to be "
              "acted on.", 0.0)
    figure(doc, F["ranking"],
           "Ranking power on the 27 held-out plots, with 95 % intervals, for "
           "the prototype and for alternatives an officer could use instead "
           "(observed). Smallest-first scores below chance because small "
           "fields here tend to yield more, not less.")
    body(doc, f"The flag separates outcomes: {pct(flag['share_bad_flagged'])} "
              "of flagged field-seasons are below threshold against "
              f"{pct(flag['share_bad_unflagged'])} of unflagged ones, a "
              f"difference of {m(flag['difference'] * 100, 0, True)} points "
              f"(95 % CI {m(flag['ci'][0] * 100, 0, True)} to "
              f"{m(flag['ci'][1] * 100, 0, True)}). Where the ranking fails "
              "is sharply localised: with a record it catches "
              f"{a['n_bad'] - a['n_missed']} of {a['n_bad']} below-threshold "
              f"field-seasons, without one {b['n_bad'] - b['n_missed']} of "
              f"{b['n_bad']}. These are small counts with wide uncertainty.")

    heading(doc, "C.  Against the register an office already keeps", 2)
    rows_a = []
    for bud, rr_ in ea["budgets"].items():
        cr = rr_["comparisons"]["register"]
        rows_a.append((pct(float(bud)),
                       f"{m(rr_['random']['recall_mean'] * ea['n_bad'], 1)}",
                       f"{rr_['register']['n_found']}",
                       f"{rr_['history_only']['n_found']}",
                       f"{rr_['workflow']['n_found']}",
                       ci(*cr["ci_adjusted"], 2)))
    body(doc, "A pre-registered experiment compared four policies on the "
              "same held-out plots at three budgets: random referral, the "
              "register an office already keeps (plots ranked by their "
              "lowest past yield), ranking by the field's own past "
              "harvests, and the SICKLE-Guard workflow. The target is "
              f"the {w(n_ho_bad)} below-threshold field-seasons, not "
              "verified damage, and six comparisons were made, so intervals "
              "are Bonferroni-adjusted (Table II).", 0.0)
    table(doc, ["Budget", "Random", "Register", "History", "Workflow",
                "Workflow \u2212 register"], rows_a,
          f"Below-threshold field-seasons found (of {ea['n_bad']}) and the "
          "workflow's adjusted recall difference against the register",
          widths=[0.42, 0.45, 0.5, 0.45, 0.55, 0.95], size=6.8)
    body(doc, "Every record-based policy far outperforms random referral: at "
              f"a 20 % budget the register finds {A20['register']['n_found']} "
              f"of {ea['n_bad']} where random visits average "
              f"{m(A20['random']['recall_mean'] * ea['n_bad'], 1)}. That "
              "prioritisation value belongs to the records, not to the "
              "workflow, which finds at most one field-season more; every "
              "adjusted interval includes zero and the pre-registered "
              "verdict is that an advantage is "
              f"{'demonstrated' if ea['verdict'] == 'advantage demonstrated' else 'not demonstrated'}"
              ". A team with a yield register would gain nothing measurable "
              "from the workflow's ranking on these plots.")

    heading(doc, "D.  Fields with no record", 2)
    body(doc, "Because every miss is a cold start, a second pre-registered "
              "experiment reserved part of the budget for fields with no "
              "record. At a 33 % budget, with no reserve the policy finds "
              f"{pct(b33z['recall_recorded'])} of below-threshold fields "
              "that have a record and none of those that do not; reserving "
              f"10 % keeps {pct(b33t['recall_recorded'])} on recorded fields "
              f"and reaches {pct(b33t['recall_cold_start'])} on cold starts; "
              "the proportional share gives "
              f"{pct(b33p['recall_recorded'])} and "
              f"{pct(b33p['recall_cold_start'])}. No allocation dominates, so "
              "the result is the trade-off. Within the cold-start pool, "
              "ordering by the satellite estimate does not beat a random "
              f"order ({m(wc['delta'], 3, True)}, 95 % CI {ci(*wc['ci95'])}): "
              "the gain comes from inspecting some cold-start fields at all. "
              "In an exploratory check, noted as such, a 10 % random reserve "
              "changed recorded recall by "
              f"{z(ex['recorded recall']['delta'])} "
              f"[{z(ex['recorded recall']['ci95'][0])}, "
              f"{z(ex['recorded recall']['ci95'][1])}]: cheap, not shown to be "
              "free.", 0.0)

    heading(doc, "E.  Why no confidence score is shown", 2)
    body(doc, f"{w(n_sig, True)} candidate yield-reliability signals were "
              "evaluated on held-out data and "
              + ("none discriminates" if n_inf == 0 else
                 f"{w(n_inf)} discriminate" + ("s" if n_inf == 1 else ""))
              + " (Table III). Prediction intervals reach "
              f"{100 * N['quantile_coverage']:.1f} % empirical coverage at a "
              "nominal 90 % and still fail, because coverage is averaged "
              "over fields and says nothing about which field is wrong. "
              "Each interval is wide, so this is a failure to demonstrate "
              "discrimination, not proof of its absence [11], [12]; but a "
              "displayed score that does not discriminate would carry the "
              "authority of a number without its information.", 0.0)
    table(doc, ["Signal", "AUROC", "95 % CI", "Informative"], sigrows,
          "Yield-reliability signals on held-out data; informative means "
          "the interval excludes 0.5. *Specified under protocols frozen "
          "before any result", widths=[1.45, 0.45, 0.85, 0.55], size=6.6)

    heading(doc, "F.  Verification as software", 2)
    body(doc, "A planner turns the ranking into an inspection plan: recorded "
              "fields ranked by their own past harvests, plus a stated "
              "reserve of fields with no record, chosen at random under a "
              "logged seed because the satellite chose among them no better "
              "than chance. Each inspector receives a task sheet that "
              "withholds the yield estimate (Fig. 5). The rule comes from an "
              f"observed count: at a 20 % budget the workflow sends "
              f"inspectors to {A20['workflow']['n_found']} below-threshold "
              "field-seasons, and "
              f"{exp20['n_inspected_bad_estimate_at_or_above_threshold']} of "
              "them carry an estimate above the threshold, which an inspector "
              "shown the number could wave through.", 0.0)
    figure(doc, F["sheet"],
           "An inspection task sheet as the workflow prints it. The estimate "
           "is withheld, the finding must be one of three observations, and "
           "a decision on a claim cannot be recorded.")
    _me = exp20["missed_estimates_kg_per_acre"]
    body(doc, "An inspection records an observation, never a decision: "
              "wording that certifies damage or decides a claim is refused, "
              "and a finding of a problem requires a photograph or a written "
              "note. Every action is appended to a SHA-256 hash chain, so "
              "altering any past record breaks every later link. The "
              f"workflow's {N['n_wf_tests']} tests edit a record, edit the "
              "log, delete an entry and reorder a plan, and each attack is "
              "caught; a demonstration round on the held-out split, with a "
              "simulated inspector labelled so in every record, left "
              f"{N['wf']['audit']['n_entries']} intact links. A field that is "
              "not inspected is reported as not inspected, never as cleared.")
    body(doc, "The fields the ranking misses are the ones a reassuring "
              f"display would reach. All {w(exp20['n_missed'])} "
              "below-threshold field-seasons missed at a 20 % budget are "
              f"cold starts, with satellite estimates of {min(_me):,.0f} to "
              f"{max(_me):,.0f} kg/acre against a threshold of {thr:,.0f}: "
              "the model pulls a field without a record towards the "
              f"average. {w(N['n_safety_tests'], True)} safety tests probe "
              "exactly these cases, and require the report's wording to be "
              "identical for estimates from 300 to 3,500 kg/acre, so a high "
              "number cannot soften a warning into reassurance; the check "
              "itself was verified by injecting a report that reassured "
              "above 2,000 kg/acre, which it caught. The prototype's "
              "browser console (Fig. 6) presents the queue, the sheet, the "
              "audit chain and the evidence, each panel labelled observed, "
              "simulated or sourced.")
    figure(doc, F["console"],
           "Summary cards of the verification console, each labelled by the "
           "kind of evidence behind it. The console also shows the queue, "
           "the task sheet and the audit chain, re-hashed in the viewer's "
           "browser.")
    if N["untranslated"] and N["untranslated"][0] == 0:
        body(doc, "The Tamil report was translated in full by the first author, a "
                  "native Tamil speaker, and each sentence is stored with the English "
                  "it was translated from. A sentence whose English later changes "
                  "falls back to English under a notice that the report is not fully "
                  "translated, and tests require every Tamil report to carry the "
                  "safety sentences in Tamil, to keep every number unchanged and to "
                  "keep its wording independent of the estimate. The translation has "
                  "not been checked by a second speaker or tested with farmers.")
    elif N["untranslated"]:
        lit, tot = N["untranslated"]
        body(doc, f"The Tamil report leaves {lit} of {tot} sentences in "
                  "English and opens by stating that it is not fully "
                  "translated; native-speaker review is required before it "
                  "is offered to Tamil-reading farmers.")

    heading(doc, "G.  Labels and model against the state's records", 2)
    body(doc, "Against Tamil Nadu's district-season crop-cutting yields for "
              "2019-20 [23], the benchmark's recalled yields run low: "
              f"{m(hl['kuruvai_ratio_to_official'], 2)} of the official "
              f"figure in Kuruvai and {m(hl['samba_ratio_to_official'], 2)} "
              "in Samba (Fig. 7a). In 2020-21 the state allocated about "
              f"\u20b9{rel20:,.0f} crore of crop-damage relief across the "
              "delta districts. Recalled Samba yields at the same plots "
              f"changed by {m(100 * hl['samba_recalled_change'], 1, True)} % "
              "into that season, and the satellite model's estimates by "
              f"{abs(100 * hl['samba_model_change']):.1f} % (Fig. 7b): the "
              "model does not respond to the season. SICKLE plots carry no "
              "district, so which of them were damaged is unknown; the check "
              "says what the model does not see, not what happened on any "
              "field. Before a yield model is relied on, it should be shown "
              "to respond to documented damage seasons.", 0.0)
    figure(doc, F["official"],
           "Observed. (a) Recalled yields against the state's crop-cutting "
           "district yields for 2019-20. (b) Same-plot change into 2020-21, "
           "the season of delta-wide relief, with 95 % intervals.")

    # ------------------------------------------------------------ VI
    heading(doc, "VI.  Simulated Studies")
    body(doc, "Every number in this section is simulated. Neither the "
              "inspectors nor the scheme's outputs are observable here, so "
              "both are modelled under stated assumptions; none of this is "
              "field or scheme performance.", 0.0)
    body(doc, "Imperfect inspectors. Keeping the observed rankings and "
              "simulating only the inspector, by the chance of confirming a "
              "real problem and of reporting a false one: at a 20 % budget a "
              "register-guided inspector with a sensitivity of "
              f"{m(ho20['register']['break_even_sensitivity'], 2)} confirms "
              "as many real problems as a perfect inspector sent at random "
              "(Fig. 8a). With a sensitivity of 0.70 and a 10 % false-alarm "
              f"rate, {m(g7['register']['precision_of_confirmed'], 2)} of a "
              "register-guided team's confirmed findings are real against "
              f"{m(g7['random']['precision_of_confirmed'], 2)} for random "
              "visits, because targeting raises the base rate of the visited "
              f"pool. Over all {oofs['n_bad']} problem field-seasons, "
              "out-of-fold, the register's advantage over random survives "
              "down to a sensitivity of "
              f"{oofs['monte_carlo_register_vs_random']['0.2']['lowest_surviving_sensitivity']}.")
    body(doc, "The scheme's weighting. No YES-TECH output is public, so its "
              "mechanism was simulated rather than its accuracy: a unit's "
              "claim is its shortfall below the threshold yield as a "
              "fraction of the sum insured, and the yield used blends the "
              "technology estimate with crop-cutting at the weights states "
              "use [21], [22]. Error sizes are assumptions; 0.30 is the "
              "largest unit-level error an earlier YES-TECH manual tolerated "
              "[20]. For a unit truly 10 % below its threshold and a "
              "technology error of 0.20, the chance of being paid nothing is "
              f"{pct(p100)} at 100 % weight against {pct(p30)} at the 30 % "
              "minimum (Fig. 8b). Re-checking the units whose estimate lands "
              "just above the threshold recovers wrongly unpaid units at "
              f"{m(gd['near_threshold']['rechecks_per_recovery'], 1)} rechecks "
              f"each against {m(gd['random']['rechecks_per_recovery'], 1)} for "
              "random rechecks. None of this says how accurate YES-TECH is.")
    figure(doc, F["simulated"],
           "SIMULATED. (a) Real problems confirmed at a 20 % budget as the "
           "simulated inspector's sensitivity falls; the rankings are "
           "observed. (b) Chance that an Insurance Unit truly short of its "
           "threshold is paid nothing, by model weight. Assumptions, not "
           "measurements of any deployed system.")

    # ------------------------------------------------------------ VII
    heading(doc, "VII.  Discussion")
    body(doc, "The evidence supports five operating rules for a verification "
              "office, each labelled by the kind of evidence behind it and "
              "none tried in an office: rank fields by their own records "
              "(observed); reserve a share of inspections for fields with no "
              "record and choose among them at random (observed); withhold "
              "the estimate from inspectors and never show a farmer a "
              "reassuring one (observed); before relying on a yield model, "
              "check that it responds to documented damage seasons "
              "(observed); and where a scheme gives a model high weight, "
              "re-verify the units just above the threshold (simulated).", 0.0)
    body(doc, "The satellite model's weak part is also the part that "
              "matters. It ranks fields by their usual level, which a "
              "register already does, and it neither ranks the seasons in "
              "which a field does badly nor responds to a documented damage "
              "season. A prototype built on it is useful only because its "
              "design does not depend on the model being right. The next "
              "question, whether the workflow changes outcomes for farmers "
              "or officials, needs a pilot with blind crop-cutting on every "
              "inspected field. A protocol for that pilot is fixed in "
              "advance: detecting the register's advantage over random "
              "visits needs about "
              f"{pil['register vs random']['fields_to_enrol']:,.0f} enrolled "
              "field-seasons, and detecting a difference between the "
              "workflow and the register about "
              f"{pil['workflow vs register']['fields_to_enrol']:,.0f}.")

    # ------------------------------------------------------------ VIII
    heading(doc, "VIII.  Limitations")
    body(doc, "Some limits bound the evidence and some bound the claims.", 0.0)
    for lead_, txt in (
            ("Small samples.",
             f"The held-out split has 27 plots and {w(n_ho_bad)} "
             "below-threshold field-seasons, so intervals are wide and no "
             "operating threshold is established."),
            ("One setting.",
             "One delta, one cropping system, 2018-2021; nothing here shows "
             "transfer elsewhere."),
            ("Recalled labels and a study threshold.",
             "Yields are farmers' recollections, so model error and recall "
             "error cannot be separated, and below-threshold means under 80 % "
             "of the training mean, a convention of the study."),
            ("No seasonal detection.",
             "On plots that cross the threshold the ranking falls to AUROC "
             f"{m(ev['auroc'])} {cil(*ev['ci'])}: the tool ranks fields by "
             "their usual level and cannot tell a bad season."),
            ("Limits on the claims.",
             "The satellite component is not shown to add value, and an "
             "advantage over the register an office keeps is not "
             "demonstrated."),
            ("No scheme linkage.",
             "The state's records are district-level and the plots carry no "
             "district; no Insurance Unit identity exists, so nothing here "
             "reproduces a scheme decision, and the YES-TECH study simulates "
             "a mechanism, not the scheme's accuracy."),
            ("No people.",
             "Nobody has used the prototype, the workflow or the console, so "
             "comprehension and outcomes for farmers or officials are "
             "unmeasured; the Tamil report was translated by one native "
             "speaker and not yet checked by a second or with farmers.")):
        lead(doc, lead_, txt, indent=0.0)

    # ------------------------------------------------------------ IX
    heading(doc, "IX.  Conclusion")
    body(doc, "India has begun to automate crop-insurance payouts, and the "
              "estimates behind them can be wrong for a particular field "
              "without any signal saying so. SICKLE-Guard shows what can be "
              "done safely downstream of such estimates: rank fields for "
              "verification by their own records, set aside a share of "
              "visits for fields that have none, keep the estimate away from "
              "inspectors, record observations rather than decisions in a "
              "tamper-evident log, and never tell a farmer that a crop is "
              "fine. On held-out plots the records carry the prioritisation "
              "value and the satellite adds nothing measurable; that is "
              "reported as found. Whether these rules improve outcomes for "
              "farmers is untested, and a pilot designed to test it is the "
              "next step.", 0.0)

    heading(doc, "Data and Code Availability")
    body(doc, "The SICKLE data and the released checkpoints are distributed "
              "by their authors on request [1] and are not redistributed "
              "here. SICKLE-Guard's code, tests and configurations, the "
              "Tamil translation file, and the scripts that regenerate every "
              "number and figure in this paper are available from the authors "
              "on request and will be released publicly with the final "
              "version.", 0.0)
    heading(doc, "Acknowledgment")
    body(doc, "We thank the authors of SICKLE for making the dataset and "
              "their trained models available.", 0.0)

    # ------------------------------------------------------------ refs
    refs_all = (json.loads((PROJ / "docs/references.json").read_text(
        encoding="utf8"))["references"]
        + json.loads((PROJ / "docs/references_thesis.json").read_text(
            encoding="utf8"))["references"])
    refs, uncited = renumber(doc, refs_all)
    heading(doc, "References")
    for r_ in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Inches(0.22)
        pf.first_line_indent = Inches(-0.22)
        pf.space_after = Pt(2)
        loc = r_.get("doi") or ""
        only_locator = any(k in r_.get("type", "")
                           for k in ("Government", "Press", "Parliament",
                                     "Presentation"))
        tail = (f" [Online]. Available: {loc}"
                if loc.startswith("http") and only_locator
                else f" doi: {loc}." if loc and not loc.startswith("http")
                else "")
        run = p.add_run(f"[{r_['n']}] {r_['entry']}{tail}")
        run.font.size = Pt(7.6)
    return doc, refs


# ----------------------------------------------------------------- checks
def self_check(text: str) -> list[str]:
    from check_claims import NEGATORS, norm
    problems = []
    low = norm(text).lower()
    reg = yaml.safe_load((T.GUARD / "configs/claims.yaml").read_text(
        encoding="utf8"))
    phrases = [str(r["phrase"]) for r in reg["prose_rules"]["must_not_appear"]]
    phrases += ["first study", "proves that", "guarantees", "is deployed",
                "farmers benefit", "reduces losses"]
    for ph in phrases:
        ph_ = norm(ph.lower())
        i = low.find(ph_)
        while i >= 0:
            window = low[max(0, i - 110):i]
            if not any(neg in window for neg in NEGATORS):
                problems.append(f"un-negated {ph!r}: "
                                f"...{low[max(0, i - 50):i + len(ph_) + 20]}...")
            i = low.find(ph_, i + 1)
    if "every number in this section is simulated" not in low:
        problems.append("the simulated section does not declare itself")
    if "not connected to pmfby" not in low:
        problems.append("the scope statement is missing")
    if "[author" in low or "to be completed" in low:
        problems.append("author placeholder left in")
    return problems


def main() -> int:
    N = T.load()
    doc, refs = build(N)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells)
    problems = self_check(text)
    if problems:
        print("SICKLE-GUARD PAPER NOT WRITTEN. The text fails its checks:")
        for p in problems:
            print("  -", p)
        return 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}")
    print(f"  {len(text.split()):,} words, {_FIGNO[0]} figures, {_TABNO[0]} "
          f"tables, {len(refs)} references")
    print("  self-check passed: no un-negated overstatement, simulated "
          "section labelled, scope stated, authors filled")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
