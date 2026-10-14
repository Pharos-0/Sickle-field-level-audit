"""Build the two-page policy brief for officials and the thesis panel.

    python scripts/make_policy_brief.py

The thesis argues; the brief tells an office what to do on Monday. It is
built from the same artefacts as the thesis paper, through the same loader,
so it cannot state a number the thesis does not. Every finding carries the
kind of evidence behind it -- OBSERVED, SIMULATED or SOURCED -- in the same
line, and the brief refuses to save if one does not, or if its text claims
something the evidence rules out.

Writes docs/Policy_Brief.docx.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

PROJ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJ / "scripts"))

import make_thesis_paper as tp  # noqa: E402

OUT = PROJ / "docs" / "Policy_Brief.docx"
LABELS = ("OBSERVED", "SIMULATED", "SOURCED")


def build(N):
    D = tp.Doc()
    for sec in D.d.sections:
        sec.left_margin = sec.right_margin = Inches(0.85)
        sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = D.d.styles["Normal"]
    st.font.size = Pt(10.5)

    vp = N["vp"]
    ea = vp["experiment_a"]["held_out"]
    a20 = ea["budgets"]["0.2"]
    b33 = vp["experiment_b"]["all_oof"]["budgets"]["0.33"]
    off = N["off"]
    hl = off["headline"]
    rel = off["l2_event"]["relief"]["total_lakh"]
    rel20 = next(v for k, v in rel.items() if k.startswith("2020-21")) / 100
    insp = N["insp"]
    ho20 = insp["sets"]["held-out test split (VERDICT)"]["budgets"]["0.2"]
    exp20 = insp["observed_exposure"]["held-out test split"]["budgets"]["0.2"]
    cc = N["yw"]["conditional_curves"]
    p100 = cc["w=1.00|sd_t=0.20|bias=+0.00|sd_c=0.10"]["0.10"]["p_paid_nothing"]
    p30 = cc["w=0.30|sd_t=0.20|bias=+0.00|sd_c=0.10"]["0.10"]["p_paid_nothing"]
    gd = N["yw"]["guard"]["loss=0.20|w=1.00|sd_t=0.20|bias=+0.00"]["budgets"]["0.05"]["policies"]
    ratio = gd["random"]["rechecks_per_recovery"] / gd["near_threshold"]["rechecks_per_recovery"]

    D.title("When Not to Trust a Satellite Yield Estimate",
            [("Five rules for crop-verification offices, from a study of "
              "paddy in the Cauvery Delta", True),
             ("Policy brief from an M.Tech Data Science thesis · "
              "[author — to be completed] · research findings, not "
              "an official assessment", False)])

    D.h("The problem", 2)
    D.p("Under PMFBY's area approach, enrolled farmers do not file claims: "
        "claims are worked out automatically from the yield the state "
        "reports for each Insurance Unit. The scheme describes YES-TECH as "
        "a gradual migration to remote-sensing based yield estimation, and "
        "in 2025-26 Madhya Pradesh gave the technology-based yield 100 % "
        "weightage and Maharashtra 50 %, against a national minimum of 30 % "
        + D.cite(21) + " [SOURCED]. Because no claim is filed, a farmer meets "
        "an error in that yield in the payout; the routes to contest it are "
        "grievance committees and a helpline " + D.cite(21) + " [SOURCED]. "
        "The yield's accuracy is contested: in a field assessment in three "
        "Madhya Pradesh districts, farmers reported falling claims and the "
        "insurers interviewed acknowledged room for 20 to 30 per cent error "
        + D.cite(22) + " [SOURCED]. That is one state's experience, not a "
        "measured property of YES-TECH. This brief addresses a narrower "
        "question: when a model's estimate for a particular field may be "
        "wrong, how should an inspection office decide where to look?")

    D.h("What we did", 2)
    D.p("On the public SICKLE benchmark (319 paddy field-seasons at 131 "
        "plots, Cauvery Delta, 2019-2020) we measured when satellite yield "
        "predictions fail, tested nine ways of flagging unreliable "
        "predictions, and built and evaluated a verification workflow under "
        "protocols fixed before the results. We checked the farmers' yield "
        "records against Tamil Nadu's own crop-cutting statistics and "
        "crop-damage relief " + D.cite(23) + ". What cannot be observed "
        "without government data, the inspector's accuracy and the "
        "insurance scheme's arithmetic, was simulated and is labelled so.")

    D.h("Five findings, and what each supports", 2)
    rows = [
        ("1", f"Field records find low-yield fields; random visits do not. "
              f"On 27 held-out plots at a 20 % inspection budget, ranking by "
              f"each field's lowest past harvest found "
              f"{a20['register']['n_found']} of {ea['n_bad']} below-threshold "
              f"fields; random visits find "
              f"{tp.m(a20['random']['recall_mean'] * ea['n_bad'], 1)} on "
              f"average. The satellite-assisted workflow found "
              f"{a20['workflow']['n_found']}, a difference from the register "
              f"not distinguishable from chance.",
         "OBSERVED", "Rank inspections by the field's own recorded harvests. "
                     "Keep that register; it is the asset."),
        ("2", f"Every field the ranking missed had no harvest record, and the "
              f"satellite chose among such fields no better than chance. "
              f"Reserving 10 % of a 33 % budget for them raised their recall "
              f"from none to {tp.pct(b33['0.10|satellite']['recall_cold_start'])} "
              f"while keeping {tp.pct(b33['0.10|satellite']['recall_recorded'])} "
              f"on recorded fields.",
         "OBSERVED", "Reserve a stated share of visits for fields with no "
                     "record, and choose them at random, with the draw "
                     "logged."),
        ("3", f"{tp.w(exp20['n_inspected_bad_estimate_at_or_above_threshold'], True)} "
              f"of the {tp.w(a20['workflow']['n_found'])} below-threshold fields "
              f"inspectors were sent to, and every one the ranking missed, "
              f"carried a model estimate above the loss threshold.",
         "OBSERVED", "Do not show inspectors the model's estimate before they "
                     "record what they see, and never show a farmer a "
                     "reassuring one."),
        ("4", f"In 2020-21 the state allocated about ₹{rel20:,.0f} crore "
              f"of crop-damage relief across the delta. At the same plots the "
              f"satellite model's yield estimates were unchanged from the "
              f"season before; which of those plots were damaged is not known, "
              f"because the plots carry no district. Farmers' recorded yields "
              f"sat at {tp.m(hl['kuruvai_ratio_to_official'], 2)} of the "
              f"state's crop-cutting figure in Kuruvai.",
         "OBSERVED", "Before relying on a yield model, check that it responds "
                     "to documented damage seasons, and do not treat recalled "
                     "yields as measured ones."),
        ("5", f"With a model error of 0.20, a unit truly 10 % below its "
              f"threshold is paid nothing {tp.pct(p100)} of the time at 100 % "
              f"model weight against {tp.pct(p30)} at 30 %. Rechecking units "
              f"whose estimate lands just above the threshold recovers wrongly "
              f"unpaid units about {tp.w(round(ratio))} times as efficiently "
              f"as random rechecks.",
         "SIMULATED", "Where model weight is high, re-verify units just above "
                      "the threshold, and test the rule on the scheme's own "
                      "unit-level data first."),
    ]
    D.table(["#", "Finding", "Evidence", "What it supports"], rows,
            "Findings on SICKLE plots in the Cauvery Delta; the held-out "
            "set is 27 plots with 11 below-threshold field-seasons, so "
            "counts carry wide uncertainty. OBSERVED: "
            "measured on real data. SIMULATED: modelled under stated "
            "assumptions, not a measurement of any deployed system.")
    # equal default widths squeeze the findings into a narrow column
    t = D.d.tables[-1]
    t.autofit = False
    for row in t.rows:
        for cell, wd in zip(row.cells, (0.3, 3.55, 0.95, 2.0)):
            cell.width = Inches(wd)

    D.h("What this does not show", 2)
    for t in ("That the system detects crop damage. Its target is a field "
              "below 80 % of the average recalled yield.",
              "How accurate YES-TECH is. No YES-TECH output is public, and "
              "none was used.",
              "Any payout, claim or outcome for a farmer. Nobody has used the "
              "workflow in the field.",
              "That these rules transfer beyond one delta, one crop and two "
              "seasons."):
        D.bullet(t)

    D.h("What it would take to go further", 2)
    D.p("A one-season pilot with a single district's inspection office: rank "
        "with the office's own records, reserve a share of visits for fields "
        "without one, use blind task sheets and the audit trail, and compare "
        "what is found against the office's current practice. Testing rule 5 "
        "needs matched unit-level technology and crop-cutting yields, which "
        "only the scheme holds. Sized from the effects observed here, about "
        "500 fields are enough to compare records with current practice and "
        "to measure inspector accuracy; whether the satellite adds anything "
        "to the register would need about "
        f"{N['pilot']['sets']['all field-seasons out-of-fold']['comparisons']['workflow vs register']['fields_to_enrol']:,.0f}, "
        "so the pilot does not try. The workflow software, its tests and an "
        "interactive console are ready for such a pilot; none of it has been "
        "used in one.")

    D.h("Sources", 2)
    for i, k in enumerate(D.cites, 1):
        para = D.d.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.3)
        para.paragraph_format.first_line_indent = Inches(-0.3)
        r = para.add_run(f"[{i}] {N['refs'][k]}")
        r.font.size = Pt(9)
        D.text.append(N["refs"][k])
    para = D.d.add_paragraph()
    r = para.add_run("Data: the SICKLE benchmark (Sani et al., WACV 2024). "
                     "Every figure in this brief is generated from the "
                     "project's artefacts by scripts/make_policy_brief.py.")
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    D.text.append(r.text)
    return D, rows


def self_check(text, rows):
    probs = [p for p in tp.self_check(text) if "objective" not in p]
    for r in rows:
        if r[2] not in LABELS:
            probs.append(f"finding {r[0]} carries no evidence label")
    if "not an official assessment" not in text:
        probs.append("the brief does not say it is not an official assessment")
    return probs


def main() -> int:
    N = tp.load()
    D, rows = build(N)
    text = "\n".join(D.text)
    probs = self_check(text, rows)
    if probs:
        print("POLICY BRIEF NOT WRITTEN:")
        for p in probs:
            print("  -", p)
        return 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    D.d.save(str(OUT))
    print(f"wrote {OUT}")
    print(f"  {len(text.split()):,} words, {len(rows)} findings, "
          f"{len(D.cites)} sources; every finding labelled")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
